"""Rebuild literal field definitions from pinned official FEC workbooks.

Run with ``uv run --with openpyxl==3.1.5 python tools/build_fec_layouts.py``.
OpenPyXL resolves worksheet relationships; sheet IDs are not file names.
This development tool makes no requests and adds no runtime dependency.
"""

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests/fixtures/fec/layouts"
DESTINATION = ROOT / "src/spicy_docs/sources/fec/field_layouts.json"


def build() -> dict:
    sources = json.loads((FIXTURES / "sources.json").read_bytes())
    layouts = []
    for source_index, source in enumerate(sources):
        path = FIXTURES / source["file"]
        raw = path.read_bytes()
        assert len(raw) == source["bytes"] and "sha256:" + hashlib.sha256(raw).hexdigest() == source["sha256"]
        workbook = load_workbook(path, read_only=True, data_only=False)
        cached = load_workbook(path, read_only=True, data_only=True)
        try:
            specification = source["file"] == "electronic-8.5.xlsx"
            if specification:
                assert workbook["HDR"]["E7"].value == "8.5"
                # Printed names differ from tab names; keep the complete summary
                # once instead of guessing which notes belong to which form.
                source["notes"] = [
                    {"sheet": "SUMMARY OF CHANGES", "row": i, "cells": [c.value for c in row]}
                    for i, row in enumerate(workbook["SUMMARY OF CHANGES"].iter_rows(), 1)
                    if any(c.value is not None for c in row)
                ]
            family = "paper" if source["file"] == "paper-headers.xlsx" else "electronic"
            for sheet in workbook:
                if sheet.title in {"NOTE", "all versions", "Version 8.5", "SUMMARY OF CHANGES"}:
                    continue
                definitions = []
                notes = []
                for row_number, (row, cached_row) in enumerate(
                    zip(sheet.iter_rows(), cached[sheet.title].iter_rows(), strict=True), 1
                ):
                    values = [c.value for c in row]
                    while values and values[-1] is None:
                        values.pop()
                    if not values:
                        continue
                    if specification:
                        if type(values[0]) is not int:
                            notes.append({"sheet": sheet.title, "row": row_number, "cells": values})
                            continue
                        assert values[0] == len(definitions) + 1 and isinstance(values[1], str)
                        # Preserve the other literal dictionary columns as observations,
                        # including examples, rule references and form associations.
                        definitions.append({"label": values[1], "cell": row[1].coordinate, "specification": values[2:]})
                    else:
                        version, form, *labels = values
                        assert isinstance(version, str) and version.startswith("v")
                        assert isinstance(form, str) and all(v is None or isinstance(v, str) for v in labels)
                        layouts.append(
                            {
                                "family": family,
                                "version": version,
                                "form": form,
                                "source_index": source_index,
                                "sheet": sheet.title,
                                "row": row[0].row,
                                "version_cell": row[0].coordinate,
                                "form_cell": row[1].coordinate,
                                "fields": [
                                    {
                                        "label": cached_row[i + 2].value if row[i + 2].data_type == "f" else value,
                                        "cell": row[i + 2].coordinate,
                                        **({"formula": value} if row[i + 2].data_type == "f" else {}),
                                    }
                                    for i, value in enumerate(labels)
                                ],
                            }
                        )
                if specification and definitions:
                    layouts.append(
                        {
                            "family": family,
                            "version": "8.5",
                            "form": sheet.title,
                            "source_index": source_index,
                            "sheet": sheet.title,
                            "row": None,
                            "version_cell": "HDR!E7",
                            "title_cell": "A1",
                            "notes": notes,
                            "fields": definitions,
                        }
                    )
        finally:
            workbook.close()
            cached.close()
    keys = [(x["family"], x["version"], x["form"], x["row"]) for x in layouts]
    assert len(keys) == len(set(keys)), "Ambiguous source layout selection"
    return {"sources": sources, "layouts": sorted(layouts, key=lambda x: (x["family"], x["version"], x["form"]))}


if __name__ == "__main__":
    catalog = build()

    # One layout per line keeps the generated source data compact and reviewable.
    def encode(value):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))

    DESTINATION.write_text(
        '{"sources":'
        + encode(catalog["sources"])
        + ',"layouts":[\n'
        + ",\n".join(encode(row) for row in catalog["layouts"])
        + "\n]}\n"
    )
    print(f"Wrote {len(catalog['layouts'])} literal source layouts to {DESTINATION}")
